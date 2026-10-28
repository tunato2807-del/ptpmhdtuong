"""Export the maintained Markdown design documents to a reviewable Word report.

Requires python-docx==1.2.0. Diagrams must already have PNG exports beside .mmd.
This is a report based on implemented scope, not a claim of external deployment.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
doc = Document()
section = doc.sections[0]
section.page_height, section.page_width = Cm(29.7), Cm(21)
section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.4)
section.left_margin, section.right_margin = Cm(2.7), Cm(2.1)
for style in ["Normal", "Body Text", "List Bullet", "List Number"]:
    doc.styles[style].font.name = "Times New Roman"
    doc.styles[style].font.size = Pt(13)
doc.styles["Normal"].paragraph_format.line_spacing = 1.3
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
for name, size in [("Title", 24), ("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)]:
    doc.styles[name].font.name = "Times New Roman"
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor.from_string("26352D")
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.keep_with_next = True


def center(text, size=14, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    return p


center("TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG TP.HCM", 14, True)
center("KHOA CÔNG NGHỆ THÔNG TIN", 14, True)
doc.add_paragraph("\n\n")
center("BÁO CÁO PHÂN TÍCH, THIẾT KẾ VÀ TRIỂN KHAI", 16, True)
center("XÂY DỰNG HỆ THỐNG QUẢN LÝ VÀ ĐẶT LỊCH\nTẠI CHUỖI CẮT TÓC NAM", 22, True)
center("GENTLEMAN · Django / PostgreSQL", 15)
doc.add_paragraph("\n\n")
center("Sinh viên: Tô Hoàng Tuấn · MSSV: 1250080221")
center("Lớp: 12_ĐH_CNPM3 · Khóa: 2023–2027")
center("Giảng viên hướng dẫn theo tài liệu gốc: ThS. Ngô Tân Khai", 12)
doc.add_paragraph("\n")
center("Bản xây dựng lại phục vụ học tập và trình diễn", 12)
doc.add_page_break()
doc.add_heading("Lưu ý sử dụng báo cáo", level=1)
doc.add_paragraph(
    "Báo cáo được biên soạn từ tài liệu đề tài do người dùng cung cấp và mã nguồn đã xây dựng trong repository. Địa chỉ, dữ liệu khách, giá dịch vụ và lịch hẹn trong bản trình diễn là dữ liệu mẫu. Thanh toán thẻ là mô phỏng, chưa có kết nối ngân hàng hoặc triển khai công khai. Người thực hiện cần rà soát nội dung, thông tin bìa và quy định của trường trước khi nộp."
)
doc.add_heading("Mục lục", level=1)
p = doc.add_paragraph()
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
p._p.append(field)
doc.add_paragraph(
    "Trong Microsoft Word, chọn toàn bộ tài liệu và cập nhật trường để tạo mục lục và số trang theo bản dàn trang của máy. Sơ đồ PNG có bản SVG và mã Mermaid đi kèm để đọc/phóng lớn."
)


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def image(path, caption):
    if path.suffix == ".svg":
        path = path.with_suffix(".png")
    if not path.exists():
        raise FileNotFoundError(path)
    # Fit both page width and height without cropping or altering the image.
    from docx.image.image import Image

    im = Image.from_file(str(path))
    max_width, max_height = Cm(16), Cm(19)
    width = min(max_width, int(max_height * im.px_width / im.px_height))
    doc.add_picture(str(path), width=width)
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(10)


for filename in [
    "01-PHAN-TICH.md",
    "02-THIET-KE.md",
    "03-TRIEN-KHAI.md",
    "04-KIEM-THU.md",
    "05-TEN-COT-DU-LIEU.md",
    "06-GIAO-DIEN-KHACH-HANG.md",
    "07-QUAN-TRI-RIENG.md",
    "11-HOAN-THIEN-CHUC-NANG.md",
    "12-FONT-VA-HINH-ANH.md",
    "13-NGHIEP-VU-NHIEU-CHI-NHANH.md",
    "14-AGENT-SKILLS.md",
]:
    doc.add_page_break()
    lines = (DOCS / filename).read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:]
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if language != "mermaid":
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.left_indent = Cm(0.3)
                r = p.add_run("\n".join(code))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            else:
                doc.add_paragraph(
                    "Kiến trúc: Trình duyệt → View/Form → Services/ORM → PostgreSQL; sự kiện qua Publisher tới Audit và Notification Observer."
                )
            i += 1
            continue
        match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if match:
            image(DOCS / match[2], match[1])
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [plain(x.strip()) for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell or "x") for cell in cells):
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Table Grid"
                for index, row in enumerate(rows):
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        cell.text = text
                        for p in cell.paragraphs:
                            p.paragraph_format.line_spacing = 1.05
                            p.paragraph_format.space_after = Pt(4)
                            for run in p.runs:
                                run.font.size = Pt(10)
                                run.bold = index == 0
            doc.add_paragraph()
            continue
        match = re.match(r"^(#{1,3})\s+(.*)", line)
        if match:
            doc.add_heading(plain(match[2]), level=len(match[1]))
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            doc.add_paragraph(plain(re.sub(r"^\d+\.\s", "", line)), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(plain(line[2:]), style="List Bullet")
        else:
            p = doc.add_paragraph(plain(line))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("GENTLEMAN · ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
doc.settings._element.append(update)
output = DOCS / "BAO_CAO_PHAN_TICH_THIET_KE.docx"
doc.save(output)
print(
    f"Created {output.name}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.inline_shapes)} diagrams/screenshots."
)
