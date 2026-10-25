# Search, pagination, Excel/PDF and appointment changes for GENTLEMAN.
import csv
import hashlib
import json
import os
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
from django import forms
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.views import LogoutView
from django.contrib.auth.decorators import login_required
from django.core import signing
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.db.models import Q, Sum, Count, Avg
from django.http import HttpResponse, HttpResponseBadRequest, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import path, reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_GET, require_http_methods
from .models import User, Branch, Service, Employee, Shift, Appointment, Invoice, AuditLog, ACTIVE_STATUSES
from .forms import RegisterForm
from .management_views import manager_required, get_resource, resource_queryset, row_values, lock_management
from .views import SignInView, operator_required, appointments_for, branches_for
from .services import local_at, validate_slot, transition, emit, assert_operator

MAX_IMPORT = 500
MAX_EXPORT = 5000
FIELDS = {
    'dich-vu': ('id','name','description','kind','price','duration','active'),
    'chi-nhanh': ('id','name','address','phone','opens_at','closes_at','active'),
    'stylist': ('id','name','branch','specialty','bio','active'),
    'ghe': ('id','branch','name','active'),
    'ca-lam': ('id','employee','start_at','end_at'),
    'tai-khoan': ('id','username','full_name','phone','email','role','branch','is_active'),
}


def page(request, title, body_template, context=None, staff=False, status=200):
    data = dict(context or {}, title=title, staff=staff,
                base='staff/base.html' if staff else 'base.html', feature_body=body_template)
    return render(request, 'features/page.html', data, status=status)


def paginate(request, queryset):
    size = request.GET.get('per_page', '20')
    size = int(size) if size in ('10','20','50') else 20
    return Paginator(queryset, size).get_page(request.GET.get('page'))


def filtered(spec, params):
    qs = resource_queryset(spec)
    query = params.get('q', '').strip()[:100]
    if query:
        condition = Q()
        for field in spec.search:
            condition |= Q(**{field+'__icontains': query})
        if query.isdecimal() and len(query) < 19:
            condition |= Q(pk=int(query))
        qs = qs.filter(condition)
    names = {field.name for field in spec.model._meta.fields}
    active = 'is_active' if spec.model is User else 'active'
    if active in names and params.get('state') in ('0','1'):
        qs = qs.filter(**{active: params['state'] == '1'})
    role = params.get('role')
    if spec.model is User and role in User.Role.values:
        qs = qs.filter(role=role)
    return qs


@manager_required
@require_GET
def resource_list(request, resource):
    spec = get_resource(resource)
    result = paginate(request, filtered(spec, request.GET))
    rows = [{'object': obj, 'values': row_values(spec, obj)} for obj in result]
    return page(request, spec.title, 'features/resource_list.html',
                {'resource':spec, 'rows':rows, 'page_obj':result, 'nav_section':'settings'}, staff=True)


@manager_required
@require_GET
def resource_detail(request, resource, pk):
    spec = get_resource(resource)
    obj = get_object_or_404(spec.model, pk=pk)
    rows = []
    for name in FIELDS[resource]:
        field = spec.model._meta.get_field(name)
        value = getattr(obj, name)
        if isinstance(value, datetime): value = timezone.localtime(value).strftime('%d/%m/%Y %H:%M')
        if isinstance(value, bool): value = 'Có' if value else 'Không'
        if field.choices: value = getattr(obj, 'get_'+name+'_display')()
        rows.append(('Mã bản ghi' if name == 'id' else field.verbose_name, value if value is not None else '—'))
    return page(request, str(obj), 'features/resource_detail.html', {'resource':spec,'object':obj,'rows':rows,'nav_section':'settings'}, staff=True)


def headers(spec):
    return [spec.model._meta.get_field(name).column for name in FIELDS[spec.key]]


def raw_values(spec, obj):
    values = []
    for name in FIELDS[spec.key]:
        field = spec.model._meta.get_field(name)
        value = getattr(obj, field.attname)
        if isinstance(value, datetime): value = timezone.localtime(value).replace(tzinfo=None)
        values.append(value)
    return values


def spreadsheet(title, columns, rows, filename):
    book = Workbook()
    sheet = book.active
    sheet.title = 'DuLieu'
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'  # Never create formulas from user text.
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='26352D')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(45, max(16, len(str(column[0].value))+4))
    notes = book.create_sheet('HuongDan')
    for text in [title, 'Mã trống: thêm mới. Mã có sẵn: cập nhật. Không đổi tên hoặc thứ tự cột.',
                 'Điện thoại nhập dạng Text để giữ số 0 đầu. Nhóm dịch vụ: hair / care / combo.',
                 'Nhập tài khoản Excel chỉ nhận vai trò customer. Quản lý vai trò nhân viên tại giao diện riêng.',
                 'Trạng thái: TRUE/FALSE. Ngày giờ: YYYY-MM-DD HH:MM, theo giờ Việt Nam.',
                 'Tài khoản mới chưa có mật khẩu đăng nhập; quản trị đặt mật khẩu bằng màn hình Sửa tài khoản.',
                 'Không có cột mật khẩu. Mỗi file nhập tối đa 500 dòng; dùng mã chi nhánh/stylist đang có.']:
        notes.append([text])
    notes.column_dimensions['A'].width = 110
    output = BytesIO(); book.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def pdf(title, columns, rows, filename):
    candidates = [os.getenv('PDF_FONT_PATH',''), 'C:/Windows/Fonts/arial.ttf',
                  '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/System/Library/Fonts/Supplemental/Arial.ttf']
    font = next((Path(name) for name in candidates if name and Path(name).is_file()), None)
    if not font: return HttpResponseBadRequest('Chưa tìm thấy font tiếng Việt. Đặt PDF_FONT_PATH tới file font .ttf trên máy.')
    if 'GentlemanPDF' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('GentlemanPDF', str(font)))
    style = getSampleStyleSheet()['Normal']; style.fontName='GentlemanPDF'; style.fontSize=8; style.leading=12
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=25, rightMargin=25, topMargin=25, bottomMargin=25)
    def text(value):
        value=' '.join(str(value if value is not None else '').split())
        return Paragraph(escape(value), style)
    table = LongTable([[text(cell) for cell in row] for row in [columns, *rows]], repeatRows=1, splitInRow=1,
                      colWidths=[(landscape(A4)[0]-50)/len(columns)]*len(columns))
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7ebdf')),
        ('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#d6dbce')),('VALIGN',(0,0),(-1,-1),'TOP'),
        ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    doc.build([text(title), text(timezone.localtime().strftime('Xuất ngày %d/%m/%Y %H:%M')), Spacer(1,14),table])
    response=HttpResponse(output.getvalue(),content_type='application/pdf')
    response['Content-Disposition']=f'attachment; filename="{filename}.pdf"'
    return response


def download(fmt, title, columns, rows, filename):
    if fmt == 'xlsx': return spreadsheet(title, columns, rows, filename)
    if fmt == 'pdf':
        if len(rows)>1000: return HttpResponseBadRequest('PDF tối đa 1.000 dòng. Hãy thu hẹp bộ lọc hoặc xuất Excel.')
        return pdf(title, columns, rows, filename)
    if fmt != 'csv': raise Http404
    response=HttpResponse(content_type='text/csv; charset=utf-8'); response.write('\ufeff')
    response['Content-Disposition']=f'attachment; filename="{filename}.csv"'
    writer=csv.writer(response); writer.writerow(columns)
    for row in rows:
        writer.writerow(["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in row])
    return response


@manager_required
@require_GET
def resource_export(request, resource, fmt):
    spec=get_resource(resource)
    template=request.GET.get('template')=='1'
    objects=[] if template else list(filtered(spec,request.GET)[:MAX_EXPORT+1])
    if len(objects)>MAX_EXPORT: return HttpResponseBadRequest('Tối đa 5.000 dòng. Hãy thu hẹp bộ lọc.')
    if template: return spreadsheet(spec.title,headers(spec),[],resource+'-mau')
    return download(fmt,spec.title,headers(spec),[raw_values(spec,obj) for obj in objects],resource)


def read_excel(upload, spec):
    if not upload or not upload.name.lower().endswith('.xlsx') or upload.size>2*1024*1024:
        raise ValueError('Chọn file .xlsx, tối đa 2 MB.')
    try:
        data=upload.read()
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist())>100 or sum(item.file_size for item in archive.infolist())>20*1024*1024:
                raise ValueError('Nội dung giải nén quá lớn.')
        book=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            sheet=book.worksheets[0]
            source=sheet.iter_rows()
            first=next(source,None)
            if first is None or [str(c.value or '').strip() for c in first] != headers(spec):
                raise ValueError('Cột không đúng mẫu. Hãy tải file mẫu tại trang này.')
            rows=[]
            for number,cells in enumerate(source,2):
                if all(cell.value is None for cell in cells): continue
                if len(rows)>=MAX_IMPORT: raise ValueError('Mỗi file tối đa 500 dòng dữ liệu.')
                values=[]
                for cell in cells:
                    value=cell.value
                    if cell.data_type in ('f','e'): raise ValueError(f'Dòng {number}: không nhận công thức hoặc ô lỗi Excel.')
                    if isinstance(value,(datetime,date,time)): value=value.isoformat()
                    elif isinstance(value,float) and value.is_integer(): value=int(value)
                    if isinstance(value,str) and len(value)>2000: raise ValueError(f'Dòng {number}: ô dữ liệu dài quá 2.000 ký tự.')
                    values.append('' if value is None else value)
                rows.append({'line':number,'data':dict(zip(FIELDS[spec.key],values))})
            if not rows: raise ValueError('File chưa có dữ liệu từ dòng 2.')
            return rows
        finally: book.close()
    except (BadZipFile, InvalidFileException, ParseError, DefusedXmlException, KeyError, OSError, RuntimeError, NotImplementedError) as error:
        raise ValueError('File Excel không hợp lệ hoặc bị hỏng.') from error


def version(obj):
    data=[str(getattr(obj,field.attname)) for field in obj._meta.fields]
    return hashlib.sha256(json.dumps(data,ensure_ascii=False).encode()).hexdigest()


def apply_rows(spec, rows, actor, expected=None, preview=False, batch=None):
    errors=[]; result=[]; versions={}; used=set()
    with transaction.atomic():
        lock_management(actor)
        marker=f'Nhập Excel {spec.key} #{batch}' if batch else ''
        if marker and AuditLog.objects.filter(actor=actor,action=marker).exists():
            raise ValueError('Phiên nhập này đã được lưu. Không ghi dữ liệu lần thứ hai.')
        for row in rows:
            number=row['line']; data=dict(row['data']); key=data.pop('id','')
            data={name: '' if value is None else str(value) for name,value in data.items()}
            try:
                with transaction.atomic():
                    if key != '':
                        if not str(key).isdecimal() or len(str(key))>18 or int(key)<1:
                            raise ValueError('Mã không hợp lệ.')
                        key=str(int(key))
                        if key in used: raise ValueError('Mã lặp trong file.')
                        used.add(key)
                        obj=spec.model.objects.select_for_update().filter(pk=int(key)).first()
                        if obj is None: raise ValueError('Mã chưa tồn tại. Để trống mã nếu muốn thêm mới.')
                        versions[str(key)]=version(obj)
                        if expected is not None and expected.get(str(key)) != versions[str(key)]:
                            raise ValueError('Bản ghi đã thay đổi sau khi xem trước. Hãy tải lại file và xem trước lần nữa.')
                        if spec.key in ('stylist','ghe') and str(data.get('branch')) != str(obj.branch_id):
                            raise ValueError('Không được chuyển chi nhánh của stylist/ghế hiện có.')
                        if spec.model is User and obj.pk==actor.pk:
                            raise ValueError('Chỉnh tài khoản đang đăng nhập tại trang Tài khoản, không nhập qua Excel.')
                    else:
                        obj=spec.model()
                        if spec.model is User: obj.set_unusable_password()
                    if spec.model is User and (obj.role!='customer' or data.get('role')!='customer'):
                        raise ValueError('Excel chỉ nhập khách hàng (role=customer). Quản lý vai trò nhân viên tại màn hình Sửa tài khoản.')
                    for name in ('active','is_active'):
                        if name in data:
                            value=str(data[name]).strip().casefold()
                            if value not in ('true','false','1','0','có','không',''): raise ValueError('Trạng thái dùng TRUE/FALSE hoặc 1/0.')
                            data[name]=value in ('true','1','có')
                    extra={'actor':actor} if spec.model is User else {}
                    form=spec.form(data,instance=obj,**extra)
                    if spec.model is User: form.fields.pop('new_password',None)
                    if not form.is_valid():
                        raise ValueError('; '.join(f'{field}: {", ".join(items)}' for field,items in form.errors.items()))
                    saved=form.save()
                    values=raw_values(spec,saved)
                    if key=='': values[0]='Tự sinh khi lưu'
                    result.append({'line':number,'action':'Cập nhật' if key != '' else 'Thêm','values':values})
            except (ValueError,ValidationError,IntegrityError) as error:
                detail='Dữ liệu trùng hoặc xung đột.' if isinstance(error,IntegrityError) else str(error)
                errors.append(f'Dòng {number}: {detail}')
        if preview or errors:
            transaction.set_rollback(True)
        else:
            AuditLog.objects.create(actor=actor,action=marker or f'Nhập Excel {spec.singular}: {len(result)} dòng'[:100])
    return result,errors,versions


@manager_required
@require_http_methods(['GET','POST'])
def resource_import(request, resource):
    spec=get_resource(resource); result=[]; errors=[]; token=''
    if request.method=='POST':
        try:
            if request.POST.get('action')=='confirm':
                payload=signing.loads(request.POST.get('token',''),salt='gentleman-import',max_age=900)
                if payload['actor']!=request.user.pk or payload['resource']!=resource: raise ValueError('Phiên xác nhận không hợp lệ.')
                batch=hashlib.sha256(request.POST['token'].encode()).hexdigest()[:32]
                result,errors,_=apply_rows(spec,payload['rows'],request.user,expected=payload['versions'],batch=batch)
                if not errors:
                    messages.success(request,f'Đã nhập {len(result)} dòng.'); return redirect('resource_list',resource=resource)
            else:
                rows=read_excel(request.FILES.get('file'),spec)
                result,errors,versions=apply_rows(spec,rows,request.user,preview=True)
                if not errors:
                    token=signing.dumps({'actor':request.user.pk,'resource':resource,'rows':rows,'versions':versions},salt='gentleman-import',compress=True)
                    if len(token)>1_500_000: raise ValueError('Dữ liệu xem trước quá lớn. Hãy chia nhỏ file.')
        except (ValueError,signing.BadSignature) as error:
            errors.append(str(error) if isinstance(error,ValueError) else 'Phiên xem trước hết hạn hoặc không hợp lệ. Hãy nhập lại file.')
    return page(request,'Nhập Excel · '+spec.title,'features/resource_import.html',
                {'resource':spec,'rows':result,'errors':errors,'token':token,'columns':headers(spec),'nav_section':'settings'},staff=True)


class ReportFilterForm(forms.Form):
    start = forms.DateField(label='Từ ngày', required=False, widget=forms.DateInput(attrs={'type': 'date', 'min': '1900-01-01', 'max': '2100-12-31'}))
    end = forms.DateField(label='Đến ngày', required=False, widget=forms.DateInput(attrs={'type': 'date', 'min': '1900-01-01', 'max': '2100-12-31'}))
    branch = forms.ModelChoiceField(label='Chi nhánh', required=False, queryset=Branch.objects.none(), empty_label='Mọi chi nhánh được phép')
    q = forms.CharField(label='Khách, điện thoại hoặc mã hóa đơn', required=False, max_length=100)

    def __init__(self, *args, user, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['branch'].queryset = branches_for(user)

    def clean(self):
        data = super().clean()
        end = data.get('end') or timezone.localdate()
        if not 1900 <= end.year <= 2100:
            self.add_error('end', 'Chọn ngày trong khoảng 1900–2100.')
            return data
        start = data.get('start') or max(date(1900, 1, 1), end-timedelta(days=29))
        if not 1900 <= start.year <= 2100:
            self.add_error('start', 'Chọn ngày trong khoảng 1900–2100.')
        elif start > end or (end-start).days > 366:
            self.add_error('start', 'Khoảng báo cáo từ 0 đến 366 ngày; ngày bắt đầu không sau ngày kết thúc.')
        data.update(start=start, end=end)
        return data


def report_queryset(request, form=None):
    form = form if form is not None else ReportFilterForm(request.GET, user=request.user)
    if not form.is_valid():
        raise ValueError(' '.join(str(error) for errors in form.errors.values() for error in errors))
    data = form.cleaned_data
    start, end = data['start'], data['end']
    qs = Invoice.objects.filter(appointment__in=appointments_for(request.user), paid_at__date__gte=start,
        paid_at__date__lte=end).select_related('appointment__customer', 'appointment__branch', 'cashier').order_by('-paid_at', '-pk')
    if data.get('branch'):
        qs = qs.filter(appointment__branch=data['branch'])
    query = data.get('q', '')
    if query:
        condition = Q(appointment__customer__full_name__icontains=query) | Q(appointment__customer__phone__icontains=query)
        code = query.upper().removeprefix('HD')
        if code.isdecimal() and len(code) < 19:
            condition |= Q(pk=int(code))
        qs = qs.filter(condition)
    return qs, start, end


@operator_required
@require_GET
def reports(request):
    form = ReportFilterForm(request.GET, user=request.user)
    valid = form.is_valid()
    if valid:
        qs, start, end = report_queryset(request, form)
        if request.GET.get('export') == 'csv':
            return report_export(request, 'csv')
    else:
        qs, start, end = Invoice.objects.none(), None, None
    summary = qs.aggregate(amount=Sum('total'), count=Count('pk'), average=Avg('total'))
    summary['total'] = summary.pop('amount') or 0
    per_branch = qs.values('appointment__branch__name').annotate(total=Sum('total'), count=Count('pk')).order_by('-total')
    return page(request, 'Báo cáo doanh thu', 'features/reports.html', {
        'page_obj': paginate(request, qs), 'summary': summary, 'filter_form': form, 'filters_valid': valid,
        'start': start, 'end': end, 'per_branch': per_branch, 'nav_section': 'reports'}, staff=True)


@operator_required
@require_GET
def report_export(request,fmt):
    try: qs,start,end=report_queryset(request)
    except ValueError as error: return HttpResponseBadRequest(str(error))
    objects=list(qs[:MAX_EXPORT+1])
    if len(objects)>MAX_EXPORT: return HttpResponseBadRequest('Tối đa 5.000 dòng. Hãy thu hẹp bộ lọc.')
    rows=[[inv.code,inv.appointment.code,inv.appointment.branch.name,inv.appointment.customer.full_name,
           timezone.localtime(inv.paid_at).replace(tzinfo=None),inv.get_method_display(),inv.total] for inv in objects]
    return download(fmt,f'Doanh thu {start:%d/%m/%Y} – {end:%d/%m/%Y}',
                    ['mahoadon','malichhen','chinhanh','khachhang','ngaythanhtoan','phuongthuc','tongtien'],rows,'doanh-thu')


@login_required
@require_GET
def invoice_pdf(request,pk):
    appointment=get_object_or_404(appointments_for(request.user),pk=pk)
    invoice=get_object_or_404(Invoice,appointment=appointment)
    rows=[[item.name,item.duration,item.price] for item in appointment.items.all()]
    rows.extend([['Giá dịch vụ (VNĐ)', '', invoice.subtotal], [f'Giảm {invoice.points_used} điểm', '', -invoice.discount], ['Tổng thanh toán (VNĐ)','',invoice.total]])
    return pdf(f'{invoice.code} · {appointment.customer.full_name} · {appointment.branch.name}',
               ['Dịch vụ','Phút','Thành tiền (VNĐ)'],rows,invoice.code)


@require_GET
def search(request):
    kind=request.GET.get('kind','dich-vu'); query=request.GET.get('q','').strip()[:100]
    options={'dich-vu':(Service,('name','description')),'chi-nhanh':(Branch,('name','address')),'stylist':(Employee,('name','specialty','branch__name'))}
    if kind not in options: kind='dich-vu'
    model,names=options[kind]; qs=model.objects.filter(active=True)
    if model is Employee: qs=qs.filter(branch__active=True).select_related('branch')
    if query:
        condition=Q()
        for name in names: condition|=Q(**{name+'__icontains':query})
        qs=qs.filter(condition)
    result=paginate(request,qs.order_by('pk')); rows=[]
    for obj in result:
        url=reverse('booking')+'?'+('service=' if model is Service else 'branch=')+str(obj.pk if model is not Employee else obj.branch_id)
        if model is Employee: url+='&employee='+str(obj.pk)
        rows.append((str(obj),getattr(obj,'description',None) or getattr(obj,'address',None) or getattr(obj,'specialty',''),url))
    return page(request,'Tìm kiếm','features/search.html',{'rows':rows,'page_obj':result,'kind':kind})


@login_required
@require_GET
def my_appointments(request):
    qs=appointments_for(request.user).filter(customer=request.user).select_related('invoice')
    query=request.GET.get('q','').strip()[:100]
    if query:
        condition=Q(branch__name__icontains=query)|Q(employee__name__icontains=query)|Q(items__name__icontains=query)
        code=query.upper().removeprefix('LH')
        if code.isdecimal() and len(code)<19: condition|=Q(pk=int(code))
        qs=qs.filter(condition).distinct()
    status=request.GET.get('status','')
    if status in Appointment.Status.values: qs=qs.filter(status=status)
    result = paginate(request, qs.order_by('-start_at', '-pk'))
    for appointment in result:
        appointment.can_manage_booking = appointment.status == 'booked' and appointment.start_at-timezone.now() >= timedelta(hours=2)
    return page(request,'Lịch hẹn của tôi','features/appointments.html',{'page_obj':result,'statuses':Appointment.Status.choices,'notifications':request.user.notifications.order_by('-created_at','-pk')[:5]})


class CancelForm(forms.Form):
    reason=forms.CharField(label='Lý do hủy (tối đa 80 ký tự)',required=False,max_length=80,widget=forms.Textarea(attrs={'rows':3}))
    confirm=forms.BooleanField(label='Tôi xác nhận hủy lịch hẹn này.')


@login_required
@require_http_methods(['GET','POST'])
def cancel(request,pk):
    appointment=get_object_or_404(appointments_for(request.user),pk=pk)
    form=CancelForm(request.POST if request.method=='POST' else None)
    if request.method=='POST' and form.is_valid():
        try:
            with transaction.atomic():
                transition(pk=pk,actor=request.user,action='cancel')
                if form.cleaned_data['reason']:
                    AuditLog.objects.create(actor=request.user,appointment=appointment,action=('Lý do hủy: '+form.cleaned_data['reason'])[:100])
            messages.success(request,'Đã hủy lịch hẹn.'); return redirect('appointment_detail',pk=pk)
        except (ValidationError,IntegrityError) as error:
            form.add_error(None,error if isinstance(error,ValidationError) else 'Lịch vừa thay đổi. Vui lòng kiểm tra lại.')
    return page(request,'Hủy '+appointment.code,'features/cancel.html',{'form':form,'appointment':appointment})


class ChangeForm(forms.Form):
    employee=forms.ModelChoiceField(label='Stylist',queryset=Employee.objects.none())
    date=forms.DateField(label='Ngày hẹn',widget=forms.DateInput(attrs={'type':'date'}))
    time=forms.TimeField(label='Giờ còn trống',required=False,widget=forms.Select)
    note=forms.CharField(label='Ghi chú',required=False,max_length=500,widget=forms.Textarea(attrs={'rows':3}))


def change_slots(appointment,employee,day):
    branch=appointment.branch; now=timezone.now(); today=timezone.localdate()
    if not branch.active or not employee.active or employee.branch_id!=branch.pk or not today<=day<=today+timedelta(days=30): return []
    duration=sum(item.duration for item in appointment.items.all())
    shifts=list(Shift.objects.filter(employee=employee,start_at__date=day))
    busy=list(Appointment.objects.filter(Q(employee=employee)|Q(customer_id=appointment.customer_id),status__in=ACTIVE_STATUSES,start_at__date=day).exclude(pk=appointment.pk))
    cursor=local_at(day,branch.opens_at)
    if cursor.second or cursor.microsecond: cursor=cursor.replace(second=0,microsecond=0)+timedelta(minutes=1)
    result=[]
    while cursor+timedelta(minutes=duration)<=local_at(day,branch.closes_at):
        end=cursor+timedelta(minutes=duration)
        if cursor>=now and cursor.minute%15==0 and any(s.start_at<=cursor and s.end_at>=end for s in shifts) and not any(a.start_at<end and a.end_at>cursor for a in busy):
            result.append(cursor.strftime('%H:%M'))
        cursor+=timedelta(minutes=1)
    return result


@login_required
@require_http_methods(['GET','POST'])
def reschedule(request,pk):
    appointment=get_object_or_404(appointments_for(request.user),pk=pk)
    initial={'employee':appointment.employee_id,'date':timezone.localdate(appointment.start_at),'time':timezone.localtime(appointment.start_at).strftime('%H:%M'),'note':appointment.note}
    form=ChangeForm(request.POST if request.method=='POST' else None,initial=initial)
    form.fields['employee'].queryset=Employee.objects.filter(branch_id=appointment.branch_id,active=True)
    form.fields['date'].widget.attrs.update(min=str(timezone.localdate()),max=str(timezone.localdate()+timedelta(days=30)))
    employee=appointment.employee; day=initial['date']
    valid=request.method=='POST' and form.is_valid()
    if valid: employee=form.cleaned_data['employee']; day=form.cleaned_data['date']
    slots=change_slots(appointment,employee,day)
    form.fields['time'].widget.choices=[('','Chọn giờ')]+[(value,value) for value in slots]
    if valid and request.POST.get('action')=='save':
        if not form.cleaned_data['time']: form.add_error('time','Chọn khung giờ trống.')
        else:
            try:
                with transaction.atomic():
                    Branch.objects.select_for_update().get(pk=appointment.branch_id)
                    current=Appointment.objects.select_for_update().get(pk=pk)
                    actor=User.objects.get(pk=request.user.pk)
                    if not actor.is_active: raise ValidationError('Tài khoản không còn hoạt động.')
                    if actor.pk!=current.customer_id: assert_operator(actor,current.branch_id)
                    if current.status!='booked': raise ValidationError('Chỉ chỉnh sửa lịch chưa check-in.')
                    if not actor.is_operator and current.start_at-timezone.now()<timedelta(hours=2):
                        raise ValidationError('Chỉ đổi lịch trước giờ hẹn ít nhất 2 giờ.')
                    customer=User.objects.select_for_update().get(pk=current.customer_id)
                    if not customer.is_active or customer.role!='customer': raise ValidationError('Khách hàng không còn hoạt động.')
                    start=local_at(day,form.cleaned_data['time'])
                    duration=sum(item.duration for item in current.items.all())
                    # Release only this allocation inside the transaction; any failure rolls it back.
                    current.status='cancelled'; current.save(update_fields=['status'])
                    branch=Branch.objects.get(pk=current.branch_id)
                    employee=Employee.objects.get(pk=employee.pk)
                    end=validate_slot(branch,employee,start,duration)
                    if Appointment.objects.filter(customer=customer,status__in=ACTIVE_STATUSES,start_at__lt=end,end_at__gt=start).exists():
                        raise ValidationError('Khách hàng đã có lịch trùng thời gian này.')
                    current.employee=employee; current.start_at=start; current.end_at=end; current.note=form.cleaned_data['note']; current.status='booked'
                    current.save(update_fields=['employee','start_at','end_at','note','status'])
                    emit(current,actor,'Đã thay đổi lịch hẹn')
                messages.success(request,'Đã cập nhật lịch hẹn, giữ nguyên mã và giá dịch vụ.'); return redirect('appointment_detail',pk=pk)
            except (ValidationError,IntegrityError) as error:
                form.add_error(None,error if isinstance(error,ValidationError) else 'Khung giờ vừa thay đổi. Hãy xem lại giờ trống.')
    return page(request,'Đổi lịch '+appointment.code,'features/reschedule.html',{'form':form,'appointment':appointment})


@require_http_methods(['GET','POST'])
def register(request):
    if request.user.is_authenticated: return redirect('home')
    form=RegisterForm(request.POST if request.method=='POST' else None)
    target=request.POST.get('next') or request.GET.get('next') or reverse('booking')
    if not url_has_allowed_host_and_scheme(target,{request.get_host()},require_https=request.is_secure()):
        target=reverse('booking')
    if request.method=='POST' and form.is_valid():
        try:
            with transaction.atomic():
                user = form.save()
        except IntegrityError:
            form.add_error('username', 'Tên đăng nhập vừa được sử dụng. Vui lòng chọn tên khác.')
        else:
            login(request, user)
            messages.success(request, 'Đăng ký thành công. Bạn đã được đăng nhập.')
            return redirect(target)
    from django.shortcuts import render
    return render(request,'auth/register.html',{'form':form,'next':target})



urlpatterns = [
    path('dang-nhap/',SignInView.as_view(),name='login'),
    path('dang-ky/',register,name='register'),
    path('dang-xuat/',LogoutView.as_view(next_page='home'),name='logout'),
    path('tim-kiem/',search,name='feature_search'),
    path('lich-cua-toi/',my_appointments,name='my_appointments'),
    path('lich-hen/<int:pk>/huy/',cancel,name='feature_cancel'),
    path('lich-hen/<int:pk>/doi-lich/',reschedule,name='feature_reschedule'),
    path('lich-hen/<int:pk>/hoa-don/pdf/',invoice_pdf,name='feature_invoice_pdf'),
    path('quan-ly/bao-cao/',reports,name='reports'),
    path('quan-ly/bao-cao/xuat/<str:fmt>/',report_export,name='feature_report_export'),
    path('quan-ly/thiet-lap/<slug:resource>/',resource_list,name='resource_list'),
    path('quan-ly/thiet-lap/<slug:resource>/nhap/',resource_import,name='feature_import'),
    path('quan-ly/thiet-lap/<slug:resource>/xuat/<str:fmt>/',resource_export,name='feature_export'),
    path('quan-ly/thiet-lap/<slug:resource>/<int:pk>/',resource_detail,name='feature_detail'),
]
