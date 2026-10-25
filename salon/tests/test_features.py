from datetime import timedelta,time
from io import BytesIO
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from django.core import signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase,Client
from django.urls import reverse
from django.utils import timezone
from salon.models import User,Branch,Service,Employee,Shift,Appointment,Invoice,AuditLog,Notification
from salon.features import headers,apply_rows,MAX_IMPORT
from salon.management_views import get_resource
from salon.services import local_at,transition
from .test_workflows import Fixtures


class FeatureTests(Fixtures,TestCase):
    def setUp(self):
        self.make_data();self.client.force_login(self.manager)

    def url(self,name,**kwargs): return reverse('feature_'+name,kwargs=kwargs)

    def excel(self,key,rows,columns=None):
        book=Workbook();book.active.append(columns or headers(get_resource(key)))
        for row in rows:book.active.append(row)
        output=BytesIO();book.save(output)
        return SimpleUploadedFile('data.xlsx',output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def preview(self,key,rows):
        return self.client.post(self.url('import',resource=key),{'action':'preview','file':self.excel(key,rows)})

    def test_opening_time_between_quarters_still_produces_valid_slots(self):
        from salon.services import available_slots
        self.branch.opens_at=time(9,10,30);self.branch.save()
        shift=Shift.objects.get(employee=self.employee,start_at__date=self.day)
        shift.start_at=local_at(self.day,time(9,10,30));shift.save()
        slots=available_slots(self.branch,self.employee,self.day,30)
        self.assertEqual(slots[0],'09:15')
        self.assertNotIn('09:10',slots)

    def test_resource_search_filters_pagination_and_read_detail(self):
        Service.objects.bulk_create([Service(name=f'Tìm kiếm {i}',price=100000,duration=30,active=i%2==0) for i in range(25)])
        response=self.client.get(reverse('resource_list',kwargs={'resource':'dich-vu'}),{'q':'Tìm kiếm','per_page':'10','page':'2'})
        self.assertEqual(response.status_code,200);self.assertEqual(len(response.context['page_obj']),10)
        self.assertContains(response,'q=T%C3%ACm+ki%E1%BA%BFm')
        response=self.client.get(reverse('resource_list',kwargs={'resource':'dich-vu'}),{'q':'Tìm kiếm','state':'0'})
        self.assertEqual(response.context['page_obj'].paginator.count,12)
        response=self.client.get(self.url('detail',resource='dich-vu',pk=self.service.pk))
        self.assertContains(response,self.service.name)
        self.assertContains(response,'Sửa');self.assertContains(response,'Xóa')

    def test_all_resource_templates_and_excel_exports_keep_vietnamese_columns(self):
        for key in ['dich-vu','chi-nhanh','stylist','ghe','ca-lam','tai-khoan']:
            with self.subTest(key=key):
                url=self.url('export',resource=key,fmt='xlsx')
                response=self.client.get(url,{'template':'1'})
                self.assertEqual(response.status_code,200)
                book=load_workbook(BytesIO(response.content));self.assertEqual(list(book.active.values)[0],tuple(headers(get_resource(key))))
                self.assertEqual(book.active.max_row,1)
                response=self.client.get(url)
                book=load_workbook(BytesIO(response.content));self.assertGreater(book.active.max_row,1)
                self.assertNotIn('matkhau',list(book.active.values)[0])

    def test_excel_preview_commit_update_and_no_changes_before_confirmation(self):
        before=Service.objects.count()
        rows=[['','Excel mới','Mô tả','hair',120000,30,True]]
        response=self.preview('dich-vu',rows)
        self.assertEqual(response.status_code,200);self.assertEqual(Service.objects.count(),before)
        token=response.context['token'];self.assertTrue(token)
        response=self.client.post(self.url('import',resource='dich-vu'),{'action':'confirm','token':token})
        self.assertEqual(response.status_code,302)
        obj=Service.objects.get(name='Excel mới');self.assertEqual(obj.price,120000)
        rows=[[obj.pk,obj.name,'Đã đổi','hair',180000,45,True]]
        response=self.preview('dich-vu',rows);token=response.context['token']
        self.client.post(self.url('import',resource='dich-vu'),{'action':'confirm','token':token})
        obj.refresh_from_db();self.assertEqual(obj.price,180000);self.assertEqual(obj.duration,45)

    def test_excel_confirm_cannot_be_replayed_for_non_unique_employee_names(self):
        response=self.preview('stylist',[['','Stylist Replay',self.branch.pk,'Fade','',True]])
        token=response.context['token']
        self.client.post(self.url('import',resource='stylist'),{'action':'confirm','token':token})
        response=self.client.post(self.url('import',resource='stylist'),{'action':'confirm','token':token})
        self.assertContains(response,'Không ghi dữ liệu lần thứ hai')
        self.assertEqual(Employee.objects.filter(name='Stylist Replay').count(),1)

    def test_registration_retains_safe_booking_target_and_never_accepts_privilege_fields(self):
        self.client.logout()
        target=reverse('booking')+'?service='+str(self.service.pk)
        response=self.client.post(reverse('register')+'?next='+target,{'username':'signup-next','full_name':'Khách mới',
            'phone':'0901111222','email':'n@example.com','password1':'New@Account2026','password2':'New@Account2026','role':'manager'})
        self.assertEqual(response.url,target)
        self.assertEqual(User.objects.get(username='signup-next').role,'customer')

    def test_import_invalid_row_rolls_back_every_valid_row(self):
        before=Service.objects.count()
        response=self.preview('dich-vu',[['','Hợp lệ','','hair',100000,30,True],['','Không hợp lệ','','hair',-10,0,True]])
        self.assertContains(response,'Dòng 3');self.assertEqual(response.context['token'],'')
        self.assertEqual(Service.objects.count(),before)

    def test_commit_revalidates_and_rejects_stale_changed_row(self):
        response=self.preview('dich-vu',[[self.service.pk,self.service.name,'Mô tả','hair',150000,30,True]])
        token=response.context['token']
        Service.objects.filter(pk=self.service.pk).update(price=333000)
        response=self.client.post(self.url('import',resource='dich-vu'),{'action':'confirm','token':token})
        self.assertContains(response,'đã thay đổi sau khi xem trước')
        self.service.refresh_from_db();self.assertEqual(self.service.price,333000)

    def test_signed_token_bound_to_user_resource_expiry_and_signature(self):
        response=self.preview('dich-vu',[['','Signed','','hair',100000,30,True]])
        token=response.context['token']
        for value in [token+'invalid',signing.dumps({'actor':self.customer.pk,'resource':'dich-vu','rows':[],'versions':{}},salt='gentleman-import')]:
            response=self.client.post(self.url('import',resource='dich-vu'),{'action':'confirm','token':value})
            self.assertEqual(response.status_code,200);self.assertFalse(Service.objects.filter(name='Signed').exists())
        with patch('django.core.signing.time.time',return_value=timezone.now().timestamp()+1000):
            response=self.client.post(self.url('import',resource='dich-vu'),{'action':'confirm','token':token})
        self.assertContains(response,'hết hạn')

    def test_bad_files_formulas_duplicate_ids_headers_and_limits_are_rejected(self):
        url=self.url('import',resource='dich-vu')
        response=self.client.post(url,{'file':SimpleUploadedFile('fake.xlsx',b'not a zip')})
        self.assertContains(response,'không hợp lệ')
        response=self.preview('dich-vu',[['','=1+1','','hair',100000,30,True]])
        self.assertContains(response,'không nhận công thức')
        response=self.preview('dich-vu',[[self.service.pk,self.service.name,'','hair',100000,30,True],['0'+str(self.service.pk),self.service.name,'','hair',100000,30,True]])
        self.assertContains(response,'Mã lặp')
        response=self.client.post(url,{'file':self.excel('dich-vu',[],['wrong'])})
        self.assertContains(response,'Cột không đúng mẫu')
        with patch('salon.features.MAX_IMPORT',1):
            response=self.preview('dich-vu',[['','A','','hair',100000,30,True],['','B','','hair',100000,30,True]])
        self.assertContains(response,'tối đa 500')

    def test_numeric_date_time_cells_return_row_errors_instead_of_server_errors(self):
        response=self.preview('chi-nhanh',[['','Invalid hours','TPHCM','0901111222',0.375,0.875,True]])
        self.assertEqual(response.status_code,200);self.assertContains(response,'Dòng 2')
        self.assertEqual(response.context['token'],'')
        self.assertFalse(Branch.objects.filter(name='Invalid hours').exists())
        response=self.preview('ca-lam',[['',self.employee.pk,12345,67890]])
        self.assertEqual(response.status_code,200);self.assertContains(response,'Dòng 2')

    def test_new_customer_import_unusable_password_and_no_staff_escalation(self):
        response=self.preview('tai-khoan',[['','excel-customer','Khách Excel','0901111222','ex@example.com','customer','',True]])
        token=response.context['token'];self.assertTrue(token)
        self.client.post(self.url('import',resource='tai-khoan'),{'action':'confirm','token':token})
        user=User.objects.get(username='excel-customer')
        self.assertFalse(user.has_usable_password());self.assertEqual(user.role,'customer')
        response=self.preview('tai-khoan',[['','injected-manager','Quản trị','','','manager','',True]])
        self.assertContains(response,'Excel chỉ nhập khách hàng')
        self.assertFalse(User.objects.filter(username='injected-manager').exists())

    def test_all_six_resources_use_existing_business_validators(self):
        start=local_at(self.day+timedelta(days=2),time(9))
        data={
            'chi-nhanh':{'id':'','name':'Nhập branch','address':'TPHCM','phone':'0901111222','opens_at':'09:00','closes_at':'21:00','active':True},
            'stylist':{'id':'','name':'Nhập stylist','branch':self.branch.pk,'specialty':'Fade','bio':'','active':True},
            'ghe':{'id':'','branch':self.branch.pk,'name':'Ghế Excel','active':True},
            'ca-lam':{'id':'','employee':self.employee.pk,'start_at':start.isoformat(),'end_at':(start+timedelta(hours=4)).isoformat()},
        }
        for key,row in data.items():
            result,errors,_=apply_rows(get_resource(key),[{'line':2,'data':row}],self.manager)
            self.assertFalse(errors,errors);self.assertEqual(len(result),1)
        appointment=self.reserve()
        shift=Shift.objects.get(employee=self.employee,start_at__date=self.day)
        row={'id':shift.pk,'employee':self.employee.pk,'start_at':local_at(self.day,time(11)).isoformat(),'end_at':shift.end_at.isoformat()}
        result,errors,_=apply_rows(get_resource('ca-lam'),[{'line':2,'data':row}],self.manager)
        self.assertTrue(errors);shift.refresh_from_db();self.assertEqual(timezone.localtime(shift.start_at).hour,9)

    def test_export_uses_filter_and_strings_never_become_excel_formulas(self):
        obj=Service.objects.create(name='=HYPERLINK("https://example.com")',price=100000,duration=30)
        response=self.client.get(self.url('export',resource='dich-vu',fmt='xlsx'),{'q':'HYPERLINK'})
        book=load_workbook(BytesIO(response.content),data_only=False)
        self.assertEqual(book.active.max_row,2);self.assertEqual(book.active['B2'].data_type,'s')
        response=self.client.get(self.url('export',resource='dich-vu',fmt='xlsx'),{'q':'does-not-exist'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row,1)
        with patch('salon.features.MAX_EXPORT',1):
            self.assertEqual(self.client.get(self.url('export',resource='dich-vu',fmt='xlsx')).status_code,400)

    def test_pdf_long_description_splits_rows_without_truncating(self):
        import re
        self.service.description='Nội dung tiếng Việt đầy đủ '*250+'KẾT THÚC'
        self.service.save()
        response=self.client.get(self.url('export',resource='dich-vu',fmt='pdf'),{'q':self.service.name})
        self.assertEqual(response.status_code,200)
        self.assertGreater(len(re.findall(rb'/Type\s*/Page\b',response.content)),1)

    def test_resource_pdf_and_invoice_pdf_access(self):
        response=self.client.get(self.url('export',resource='dich-vu',fmt='pdf'))
        self.assertEqual(response.status_code,200);self.assertTrue(response.content.startswith(b'%PDF-'))
        a=self.reserve();a.status='completed';a.save();Invoice.objects.create(appointment=a,total=a.total,method='cash',cashier=self.manager)
        self.client.force_login(self.customer)
        self.assertTrue(self.client.get(self.url('invoice_pdf',pk=a.pk)).content.startswith(b'%PDF-'))
        self.client.force_login(self.other_customer)
        self.assertEqual(self.client.get(self.url('invoice_pdf',pk=a.pk)).status_code,404)

    def test_resource_import_export_detail_permissions_and_csrf(self):
        for user in [self.customer,self.receptionist]:
            self.client.force_login(user)
            for name,extra in [('import',{}),('export',{'fmt':'xlsx'}),('detail',{'pk':self.service.pk})]:
                self.assertEqual(self.client.get(self.url(name,resource='dich-vu',**extra)).status_code,403)
        client=Client(enforce_csrf_checks=True);client.force_login(self.manager)
        self.assertEqual(client.post(self.url('import',resource='dich-vu'),{'action':'confirm','token':'bad'}).status_code,403)

    def test_report_filters_and_exports_cannot_leak_other_branches(self):
        a=self.reserve();a.status='completed';a.save();Invoice.objects.create(appointment=a,total=a.total,method='cash',cashier=self.manager)
        b=self.reserve(customer=self.other_customer,branch_id=self.other_branch.pk,employee_id=self.other_employee.pk)
        b.status='completed';b.save();Invoice.objects.create(appointment=b,total=b.total,method='cash',cashier=self.manager)
        self.client.force_login(self.receptionist)
        response=self.client.get(self.url('report_export',fmt='xlsx'))
        rows=list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][2],self.branch.name)
        self.assertEqual(self.client.get(self.url('report_export',fmt='xlsx'),{'branch':self.other_branch.pk}).status_code,400)
        self.assertEqual(self.client.get(self.url('report_export',fmt='pdf'),{'start':'invalid'}).status_code,400)
        response=self.client.get(reverse('reports'),{'q':a.customer.full_name})
        self.assertContains(response,'Xuất Excel');self.assertContains(response,'Xuất PDF');self.assertContains(response,'Xuất CSV')

    def test_public_search_ignores_inactive_data_and_paginates(self):
        self.client.logout()
        self.service.active=False;self.service.save()
        response=self.client.get(self.url('search'),{'q':self.service.name})
        self.assertEqual(response.status_code,200);self.assertEqual(response.context['page_obj'].paginator.count,0)
        response=self.client.get(self.url('search'),{'kind':'stylist','q':'Stylist'})
        self.assertEqual(response.context['page_obj'].paginator.count,2)

    def test_staff_can_find_booking_code_without_cross_branch_leaks(self):
        own=self.reserve()
        other=self.reserve(customer=self.other_customer,branch_id=self.other_branch.pk,employee_id=self.other_employee.pk)
        self.client.force_login(self.receptionist)
        response=self.client.get(reverse('dashboard'),{'date':str(self.day),'q':own.code})
        self.assertEqual(response.context['appointments'].paginator.count,1)
        response=self.client.get(reverse('dashboard'),{'date':str(self.day),'q':other.code})
        self.assertEqual(response.context['appointments'].paginator.count,0)

    def test_personal_search_and_pagination_only_own_appointments(self):
        self.reserve();self.reserve(customer=self.other_customer,start=local_at(self.day,time(11)))
        self.client.force_login(self.customer)
        response=self.client.get(reverse('my_appointments'),{'q':self.branch.name,'status':'booked','per_page':'10'})
        self.assertEqual(response.context['page_obj'].paginator.count,1)
        self.assertContains(response,'Đổi lịch');self.assertContains(response,'Hủy lịch')
        for index in range(7): Notification.objects.create(recipient=self.customer,message='Mới '+str(index))
        response=self.client.get(reverse('my_appointments'))
        self.assertEqual(list(response.context['notifications'])[0].message,'Mới 6')
        self.assertEqual(len(response.context['notifications']),5)

    def test_cancel_confirmation_reason_and_two_hour_policy(self):
        a=self.reserve();url=self.url('cancel',pk=a.pk)
        self.client.force_login(self.other_customer);self.assertEqual(self.client.get(url).status_code,404)
        self.client.force_login(self.customer)
        self.assertEqual(self.client.get(url).status_code,200)
        a.refresh_from_db();self.assertEqual(a.status,'booked')
        self.client.post(url,{'confirm':'on','reason':'Bận việc'})
        a.refresh_from_db();self.assertEqual(a.status,'cancelled')
        self.assertTrue(AuditLog.objects.filter(appointment=a,action__contains='Bận việc').exists())
        count=AuditLog.objects.filter(appointment=a).count()
        self.client.post(url,{'confirm':'on'});self.assertEqual(AuditLog.objects.filter(appointment=a).count(),count)
        b=self.reserve(start=local_at(self.day,time(12)))
        with patch('salon.services.timezone.now',return_value=b.start_at-timedelta(minutes=60)):
            response=self.client.post(self.url('cancel',pk=b.pk),{'confirm':'on'})
        self.assertContains(response,'ít nhất 2 giờ');b.refresh_from_db();self.assertEqual(b.status,'booked')

    def test_reschedule_keeps_code_price_and_duration_and_checks_conflicts(self):
        a=self.reserve();old=a.pk
        self.service.price=999000;self.service.duration=60;self.service.save()
        self.client.force_login(self.customer)
        payload={'employee':self.employee.pk,'date':str(self.day),'time':'11:00','note':'Đã đổi','action':'save'}
        response=self.client.post(self.url('reschedule',pk=a.pk),payload)
        self.assertEqual(response.status_code,302)
        a.refresh_from_db();self.assertEqual(a.pk,old);self.assertEqual(a.total,100000)
        self.assertEqual(a.end_at-a.start_at,timedelta(minutes=30));self.assertEqual(timezone.localtime(a.start_at).hour,11)
        self.reserve(customer=self.other_customer,start=local_at(self.day,time(12)))
        payload['time']='12:00'
        response=self.client.post(self.url('reschedule',pk=a.pk),payload)
        self.assertEqual(response.status_code,200);a.refresh_from_db();self.assertEqual(a.status,'booked')
        self.assertEqual(timezone.localtime(a.start_at).hour,11)

    def test_reschedule_access_state_time_and_customer_conflicts(self):
        a=self.reserve()
        self.client.force_login(self.other_customer)
        self.assertEqual(self.client.get(self.url('reschedule',pk=a.pk)).status_code,404)
        self.client.force_login(self.customer)
        payload={'employee':self.employee.pk,'date':str(self.day),'time':'11:00','action':'save'}
        with patch('salon.features.timezone.now',return_value=a.start_at-timedelta(hours=1)):
            response=self.client.post(self.url('reschedule',pk=a.pk),payload)
        self.assertContains(response,'ít nhất 2 giờ')
        a.status='completed';a.save()
        response=self.client.post(self.url('reschedule',pk=a.pk),payload)
        self.assertContains(response,'Chỉ chỉnh sửa lịch chưa check-in')
        a.status='booked';a.save()
        self.reserve(branch_id=self.other_branch.pk,employee_id=self.other_employee.pk,start=local_at(self.day,time(11)))
        response=self.client.post(self.url('reschedule',pk=a.pk),payload)
        self.assertContains(response,'Khách hàng đã có lịch trùng')
        a.refresh_from_db();self.assertEqual(a.status,'booked');self.assertEqual(timezone.localtime(a.start_at).hour,10)

    def test_login_logout_and_invalid_redirect_still_work(self):
        self.client.logout()
        self.assertFalse(self.client.login(username='customer',password='wrong'))
        response=self.client.post(reverse('login')+'?next=https://evil.example/',{'username':'customer','password':'Test@Password26'})
        self.assertEqual(response.url,reverse('my_appointments'))
        self.assertEqual(self.client.get(reverse('logout')).status_code,405)
        self.assertEqual(self.client.post(reverse('logout')).status_code,302)
        self.assertEqual(self.client.get(reverse('my_appointments')).status_code,302)


from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from django.test import TransactionTestCase
from django.db import connections

class RescheduleConcurrencyTests(Fixtures,TransactionTestCase):
    def setUp(self): self.make_data()

    def test_two_changes_cannot_take_same_slot_and_loser_keeps_old_booking(self):
        first=self.reserve()
        second=self.reserve(customer=self.other_customer,start=local_at(self.day,time(11)))
        barrier=Barrier(2)
        def attempt(pair):
            uid,aid=pair
            try:
                client=Client();client.force_login(User.objects.get(pk=uid))
                barrier.wait(timeout=10)
                response=client.post(reverse('feature_reschedule',kwargs={'pk':aid}),
                    {'employee':self.employee.pk,'date':str(self.day),'time':'12:00','action':'save'})
                return response.status_code
            finally: connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(attempt,[(self.customer.pk,first.pk),(self.other_customer.pk,second.pk)]))
        self.assertCountEqual(results,[200,302])
        self.assertEqual(Appointment.objects.filter(start_at=local_at(self.day,time(12))).count(),1)
        self.assertEqual(Appointment.objects.filter(status='booked').count(),2)
